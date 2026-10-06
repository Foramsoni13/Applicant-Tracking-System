import sys
import json
import os
import re
import traceback

# ==========================================
# DEPENDENCY IMPORTS WITH GRACEFUL FALLBACK
# ==========================================

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None

try:
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
except ImportError:
    TfidfVectorizer = None
    cosine_similarity = None


# ==========================================
# CANONICAL ALIAS DICTIONARIES
# ==========================================

DEGREE_ALIASES = {
    "mca": ["mca", "master of computer applications", "master of computer application", "masters in computer applications"],
    "bca": ["bca", "bachelor of computer applications", "bachelor of computer application", "bachelors in computer applications"],
    "btech": ["b.tech", "btech", "b.e", "be", "bachelor of technology", "bachelor of engineering", "bachelors in technology", "bachelors of engineering"],
    "mtech": ["m.tech", "mtech", "m.e", "me", "master of technology", "master of engineering", "masters in technology", "masters of engineering"],
    "bsc": ["b.sc", "bsc", "bs", "bachelor of science", "bachelors in science"],
    "msc": ["m.sc", "msc", "ms", "master of science", "masters in science"],
    "mba": ["mba", "master of business administration", "masters in business administration"],
    "diploma": ["diploma"],
    "12th": ["12th", "hsc", "higher secondary"],
    "10th": ["10th", "ssc", "secondary school"]
}

SKILL_ALIASES = {
    "react": ["react", "react.js", "reactjs", "react js"],
    "javascript": ["javascript", "js", "java script"],
    "node.js": ["node.js", "nodejs", "node js", "node"],
    "mongodb": ["mongodb", "mongo db", "mongo"],
    "python": ["python", "py"],
    "html": ["html", "html5"],
    "css": ["css", "css3"],
    "express": ["express", "express.js", "expressjs"],
    "typescript": ["typescript", "ts"],
    "sql": ["sql", "mysql", "postgresql", "postgres", "sqlite"],
    "java": ["java", "core java", "java ee", "j2ee"],
    "c++": ["c++", "cpp"],
    "c#": ["c#", "csharp", "c sharp"],
    ".net": [".net", "asp.net", "dotnet"],
    "vue": ["vue", "vue.js", "vuejs"],
    "angular": ["angular", "angularjs", "angular.js"],
    "django": ["django"],
    "flask": ["flask"],
    "git": ["git", "github", "gitlab"],
    "docker": ["docker"],
    "aws": ["aws", "amazon web services"],
}

STOP_WORDS = {
    "a", "an", "the", "and", "or", "but", "in", "on", "at", "to", "for",
    "of", "with", "by", "from", "is", "it", "as", "was", "were", "be",
    "been", "being", "have", "has", "had", "do", "does", "did", "will",
    "would", "shall", "should", "may", "might", "can", "could", "am",
    "are", "this", "that", "these", "those", "i", "me", "my", "we",
    "our", "you", "your", "he", "him", "his", "she", "her", "they",
    "them", "their", "its", "not", "no", "so", "if", "then", "than",
    "too", "very", "just", "about", "after", "all", "also", "any",
    "between", "both", "each", "here", "how", "into", "more", "most",
    "other", "out", "over", "own", "same", "some", "such", "through",
    "under", "until", "up", "what", "when", "where", "which", "while",
    "who", "why", "etc", "eg", "ie", "vs"
}


# ==========================================
# FAST TEXT EXTRACTION
# ==========================================

def extract_text_pymupdf(file_path):
    if fitz is None:
        raise ImportError("PyMuPDF not available")
    text = ""
    doc = fitz.open(file_path)
    for page in doc:
        text += page.get_text("text") + "\n"
    doc.close()
    return text.strip()


def extract_text_pdfplumber(file_path):
    if pdfplumber is None:
        raise ImportError("pdfplumber not available")
    text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
    return text.strip()


def extract_text_docx(file_path):
    if docx is None:
        raise ImportError("python-docx not available")
    document = docx.Document(file_path)
    lines = []
    for para in document.paragraphs:
        if para.text.strip():
            lines.append(para.text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    lines.append(cell.text)
    return "\n".join(lines).strip()


def extract_text(file_path):
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()
    errors = []

    if ext == ".pdf":
        if fitz is not None:
            try:
                text = extract_text_pymupdf(file_path)
                if text and len(text.strip()) > 10:
                    return text
            except Exception as e:
                errors.append(f"PyMuPDF: {e}")

        if pdfplumber is not None:
            try:
                text = extract_text_pdfplumber(file_path)
                if text and len(text.strip()) > 10:
                    return text
            except Exception as e:
                errors.append(f"pdfplumber: {e}")

        if errors:
            raise ValueError(f"PDF parsing failed. Reason: {'; '.join(errors)}")
        raise ValueError("No PDF parser available. Install pymupdf or pdfplumber.")

    elif ext in [".docx", ".doc"]:
        try:
            text = extract_text_docx(file_path)
            if text and len(text.strip()) > 10:
                return text
            raise ValueError("DOCX is empty or could not be parsed.")
        except Exception as e:
            raise ValueError(f"DOCX parsing failed: {e}")

    else:
        raise ValueError(f"Unsupported file type: {ext}. Upload PDF or DOCX.")


# ==========================================
# METADATA EXTRACTION
# ==========================================

def extract_email(text):
    match = re.search(r'[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}', text)
    return match.group(0) if match else "Not specified"


def extract_phone(text):
    patterns = [
        r'(?:\+?\d{1,3}[\s\-]?)?\(?\d{3}\)?[\s\-]?\d{3}[\s\-]?\d{4}',
        r'\b\d{10}\b',
        r'(?:\+?\d{1,3}[\s\-]?)?\d{5}[\s\-]?\d{5}',
    ]
    for p in patterns:
        m = re.search(p, text)
        if m:
            return m.group(0).strip()
    return "Not specified"


def extract_name(text):
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    for line in lines[:5]:
        cleaned = re.sub(r'[^a-zA-Z\s.]', '', line).strip()
        if (
            3 <= len(cleaned) <= 60 and
            "@" not in cleaned and
            not cleaned[0].isdigit() and
            not re.search(
                r'resume|curriculum|vitae|cv|objective|summary|profile|phone|email|address|contact',
                cleaned, re.I
            )
        ):
            words = [w for w in cleaned.split() if len(w) > 1]
            if 2 <= len(words) <= 5:
                return " ".join(w.capitalize() for w in words)
    return "Candidate Name"


# ==========================================
# DEGREE & EDUCATION MATCHING
# ==========================================

def get_canonical_degrees(text):
    text_lower = text.lower()
    found_canonical = set()
    for canonical_key, aliases in DEGREE_ALIASES.items():
        for alias in aliases:
            pattern = r'\b' + re.escape(alias) + r'\b'
            if re.search(pattern, text_lower):
                found_canonical.add(canonical_key)
                break
    return found_canonical


def match_education(resume_text, job_education_req, job_desc, structured_education=None):
    has_edu = (structured_education and len(structured_education) > 0) or bool(get_canonical_degrees(resume_text))
    if not has_edu:
        return {
            "score": 0.0,
            "matched": False,
            "candidate_degrees": ["None"],
            "required_degrees": []
        }

    req_text = (str(job_education_req) + " " + str(job_desc)).lower()
    required_degrees = get_canonical_degrees(req_text)
    candidate_degrees = get_canonical_degrees(resume_text)

    if not required_degrees or "any education" in str(job_education_req).lower():
        return {
            "score": 100.0,
            "matched": True,
            "candidate_degrees": list(candidate_degrees) if candidate_degrees else ["Detected"],
            "required_degrees": ["Any Education"]
        }

    if candidate_degrees.intersection(required_degrees):
        return {
            "score": 100.0,
            "matched": True,
            "candidate_degrees": list(candidate_degrees),
            "required_degrees": list(required_degrees)
        }
    
    if candidate_degrees:
        return {
            "score": 40.0,
            "matched": False,
            "candidate_degrees": list(candidate_degrees),
            "required_degrees": list(required_degrees)
        }

    return {
        "score": 0.0,
        "matched": False,
        "candidate_degrees": ["Not detected"],
        "required_degrees": list(required_degrees)
    }


# ==========================================
# SKILL EXTRACTION & MATCHING
# ==========================================

def extract_all_skills_from_text(resume_text):
    """Extract skills present in resume using canonical dictionary and regex patterns."""
    text_lower = resume_text.lower()
    found_skills = []

    for canonical_key, aliases in SKILL_ALIASES.items():
        for alias in aliases:
            if alias in [".net", "c++", "c#"]:
                pattern = re.escape(alias)
            else:
                pattern = r'\b' + re.escape(alias) + r'\b'

            if re.search(pattern, text_lower):
                found_skills.append(aliases[0].capitalize() if len(aliases[0]) > 3 else aliases[0].upper())
                break
    return sorted(list(set(found_skills)))


def match_skills_strictly(resume_text, job_skills):
    extracted = extract_all_skills_from_text(resume_text)
    if not extracted:
        return {
            "score": 0.0,
            "matched": [],
            "missing": job_skills or [],
            "extra": []
        }

    if not job_skills or len(job_skills) == 0:
        score = min((len(extracted) / 5) * 100.0, 100.0)
        return {
            "score": score,
            "matched": extracted,
            "missing": [],
            "extra": []
        }

    text_lower = resume_text.lower()
    matched = []
    missing = []

    for jskill in job_skills:
        jskill_clean = jskill.strip()
        jskill_lower = jskill_clean.lower()

        # Find aliases if available
        aliases = [jskill_lower]
        for canonical_key, alias_list in SKILL_ALIASES.items():
            if jskill_lower in alias_list or canonical_key == jskill_lower:
                aliases = alias_list
                break

        is_found = False
        for alias in aliases:
            if alias in [".net", "c++", "c#"]:
                pattern = re.escape(alias)
            else:
                pattern = r'\b' + re.escape(alias) + r'\b'

            if re.search(pattern, text_lower):
                is_found = True
                break

        if is_found:
            matched.append(jskill_clean)
        else:
            missing.append(jskill_clean)

    score = (len(matched) / len(job_skills)) * 100.0
    extra = [s for s in extracted if s.lower() not in [js.lower() for js in job_skills]]

    return {
        "score": score,
        "matched": matched,
        "missing": missing,
        "extra": extra
    }


# ==========================================
# EXPERIENCE MATCHING
# ==========================================

def extract_experience_years(text):
    text_lower = str(text).lower()
    if "fresher" in text_lower or "no experience" in text_lower or "0 years" in text_lower:
        return 0.0
    matches = re.findall(r'(\d+(?:\.\d+)?)\s*\+?\s*(?:year|yr)s?', text_lower)
    if matches:
        return max([float(m) for m in matches])
    return None


def match_experience(resume_text, job_experience, structured_experience=None):
    has_exp = (structured_experience and len(structured_experience) > 0) or bool(re.search(r'experience|work history|employment history|internship', resume_text, re.I))

    # IF RESUME HAS NO EXPERIENCE AT ALL -> 0.0%
    if not has_exp:
        return 0.0

    req_years = extract_experience_years(job_experience)
    cand_years = extract_experience_years(resume_text)

    if req_years is None or req_years == 0:
        return 80.0 if has_exp else 0.0

    if cand_years is None or cand_years == 0:
        return 40.0 if has_exp else 0.0

    if cand_years >= req_years:
        return 100.0
    else:
        return round(min((cand_years / req_years) * 100.0, 100.0), 1)


# ==========================================
# STRUCTURED SECTIONS EXTRACTION
# ==========================================

def extract_section_block(resume_text, header_patterns, next_header_patterns):
    lines = resume_text.split("\n")
    start_idx = -1
    for i, line in enumerate(lines):
        clean_l = line.strip().lower()
        for hp in header_patterns:
            if re.search(hp, clean_l):
                start_idx = i + 1
                break
        if start_idx != -1:
            break

    if start_idx == -1:
        return []

    block_lines = []
    for line in lines[start_idx:]:
        clean_l = line.strip().lower()
        is_next_header = False
        for nhp in next_header_patterns:
            if re.search(nhp, clean_l):
                is_next_header = True
                break
        if is_next_header:
            break
        block_lines.append(line)

    return [l.strip() for l in block_lines if l.strip()]


def extract_education_structured(resume_text):
    header_patterns = [
        r'^\s*(?:academic\s+|educational\s+)?education\s*(?:qualification|background)?\s*$',
        r'^\s*qualifications?\s*$', r'^\s*academic\s+background\s*$'
    ]
    next_header_patterns = [
        r'^\s*(?:work\s+)?experience\s*$', r'^\s*projects?\s*$', r'^\s*skills?\s*$',
        r'^\s*certificat(?:ions?|es?)\s*$', r'^\s*achievements?\s*$', r'^\s*languages?\s*$'
    ]

    block_lines = extract_section_block(resume_text, header_patterns, next_header_patterns)
    if not block_lines:
        lines = [l.strip() for l in resume_text.split("\n") if l.strip()]
        for line in lines:
            if get_canonical_degrees(line) or re.search(r'university|college|institute|school|academy', line, re.I):
                block_lines.append(line)

    results = []
    seen = set()
    current = {}

    for line in block_lines:
        clean_l = re.sub(r'^[•\-\*\d\.\s]+', '', line).strip()
        degrees = get_canonical_degrees(clean_l)
        
        # Check if line contains a degree
        if degrees and not current.get("degree"):
            current["degree"] = clean_l
            continue

        # Check if line contains institution
        if re.search(r'university|college|institute|school|academy', clean_l, re.I) and not current.get("institution"):
            current["institution"] = clean_l
            continue

        # Check if line contains year
        year_match = re.search(r'\b(19\d{2}|20\d{2})\b', clean_l)
        if year_match and not current.get("year"):
            current["year"] = year_match.group(0)

        # Check if line contains CGPA / percentage
        cgpa_match = re.search(r'(\d{1,2}(?:\.\d+)?(?:\s*[\/%]\s*10)?\s*(?:cgpa|gpa|%)?)', clean_l, re.I)
        if cgpa_match and re.search(r'cgpa|gpa|%', clean_l, re.I) and not current.get("cgpa"):
            current["cgpa"] = cgpa_match.group(0)

        if current.get("degree") or current.get("institution"):
            if len(current.keys()) >= 2 or len(clean_l) > 60:
                deg = current.get("degree", "Degree")
                inst = current.get("institution", "")
                yr = current.get("year", "")
                cg = current.get("cgpa", "")
                key = f"{deg.lower()}|{inst.lower()}"
                if key not in seen:
                    seen.add(key)
                    results.append({"degree": deg, "institution": inst, "year": yr, "cgpa": cg})
                current = {}

    if current.get("degree") or current.get("institution"):
        deg = current.get("degree", "Degree")
        inst = current.get("institution", "")
        yr = current.get("year", "")
        cg = current.get("cgpa", "")
        key = f"{deg.lower()}|{inst.lower()}"
        if key not in seen:
            seen.add(key)
            results.append({"degree": deg, "institution": inst, "year": yr, "cgpa": cg})

    if not results:
        degrees = get_canonical_degrees(resume_text)
        for d in degrees:
            d_display = d.upper() if len(d) <= 5 else d.capitalize()
            inst = ""
            for line in resume_text.split("\n"):
                if re.search(r'university|college|institute|school', line, re.I):
                    inst = line.strip()
                    break
            results.append({"degree": d_display, "institution": inst, "year": "", "cgpa": ""})

    return results


def extract_certifications_structured(resume_text):
    header_patterns = [
        r'^\s*(?:professional\s+)?certificat(?:ions?|es?)\s*$',
        r'^\s*licenses?\s+(?:and|&)\s+certifications?\s*$'
    ]
    next_header_patterns = [
        r'^\s*(?:work\s+)?experience\s*$', r'^\s*projects?\s*$', r'^\s*education\s*$',
        r'^\s*skills?\s*$', r'^\s*achievements?\s*$', r'^\s*languages?\s*$'
    ]
    block_lines = extract_section_block(resume_text, header_patterns, next_header_patterns)
    if not block_lines:
        lines = [l.strip() for l in resume_text.split("\n") if l.strip()]
        for line in lines:
            if re.search(r'\bcertif(?:ied|icate|ication)\b', line, re.I) and len(line) < 120:
                block_lines.append(line)

    results = []
    seen = set()
    for line in block_lines:
        clean_l = re.sub(r'^[•\-\*\d\.\s]+', '', line).strip()
        if 3 <= len(clean_l) <= 120 and clean_l.lower() not in seen and not re.search(r'^\s*certificat(?:ions?|es?)\s*$', clean_l, re.I):
            seen.add(clean_l.lower())
            results.append(clean_l)
    return results


def extract_languages_structured(resume_text):
    known = ["english", "hindi", "gujarati", "marathi", "spanish", "french", "german", "mandarin", "chinese", "japanese", "korean", "russian", "arabic", "portuguese", "italian", "bengali", "tamil", "telugu", "kannada", "malayalam", "punjabi", "urdu"]
    lower = resume_text.lower()
    found = set()
    for lang in known:
        if re.search(r'\b' + lang + r'\b', lower):
            found.add(lang.capitalize())
    return sorted(list(found))


def extract_achievements_structured(resume_text):
    header_patterns = [
        r'^\s*achievements?\s*$', r'^\s*awards?\s*(?:and|&)\s*achievements?\s*$', r'^\s*honors?\s*$'
    ]
    next_header_patterns = [
        r'^\s*(?:work\s+)?experience\s*$', r'^\s*projects?\s*$', r'^\s*education\s*$',
        r'^\s*skills?\s*$', r'^\s*certificat(?:ions?|es?)\s*$', r'^\s*languages?\s*$'
    ]
    block_lines = extract_section_block(resume_text, header_patterns, next_header_patterns)
    results = []
    seen = set()
    for line in block_lines:
        clean_l = re.sub(r'^[•\-\*\d\.\s]+', '', line).strip()
        if 5 <= len(clean_l) <= 150 and clean_l.lower() not in seen:
            seen.add(clean_l.lower())
            results.append(clean_l)
    return results


def match_projects(resume_text, structured_projects, job_skills, job_desc):
    if not structured_projects:
        return 0.0

    job_req_text = ((" ".join(job_skills)) + " " + str(job_desc)).lower().strip()
    if not job_req_text:
        return 100.0

    matched_proj_count = 0
    for proj in structured_projects:
        proj_text = (proj.get("name", "") + " " + proj.get("description", "") + " " + (" ".join(proj.get("technologies", [])))).lower()
        has_match = False
        if job_skills:
            for js in job_skills:
                if js.lower() in proj_text:
                    has_match = True
                    break
        else:
            words = [w for w in job_req_text.split() if len(w) > 3 and w not in STOP_WORDS]
            for w in words[:10]:
                if w in proj_text:
                    has_match = True
                    break
        if has_match:
            matched_proj_count += 1

    if len(structured_projects) > 0:
        ratio = matched_proj_count / len(structured_projects)
        if ratio > 0:
            return round(min(50.0 + (ratio * 50.0), 100.0), 1)
        return 50.0
    return 0.0


def match_certifications(resume_text, structured_certs, job_desc):
    has_certs = (structured_certs and len(structured_certs) > 0) or bool(re.search(r'certif', resume_text, re.I))

    # IF RESUME HAS NO CERTIFICATIONS AT ALL -> 0.0%
    if not has_certs:
        return 0.0

    return 100.0


def calculate_completeness_score(candidate_name, email, phone, location, skills, education, experience, projects, certifications):
    score = 0.0

    if candidate_name and candidate_name != "Candidate Name":
        score += 15.0
    if email and email != "Not specified":
        score += 15.0
    if phone and phone != "Not specified":
        score += 15.0
    if location and location != "Not specified":
        score += 10.0
    if skills and len(skills) > 0:
        score += 15.0
    if education and len(education) > 0:
        score += 10.0
    if experience and len(experience) > 0:
        score += 10.0
    if projects and len(projects) > 0:
        score += 10.0

    return round(min(score, 100.0))


def extract_experience_structured(resume_text):
    header_patterns = [
        r'^\s*(?:work\s+|professional\s+|employment\s+)?experience\s*$',
        r'^\s*work\s+history\s*$', r'^\s*employment\s+history\s*$', r'^\s*internships?\s*$'
    ]
    next_header_patterns = [
        r'^\s*projects?\s*$', r'^\s*education\s*$', r'^\s*skills?\s*$',
        r'^\s*certificat(?:ions?|es?)\s*$', r'^\s*declarat(?:ion|ions)\s*$',
        r'^\s*achievements?\s*$', r'^\s*languages?\s*$'
    ]

    block_lines = extract_section_block(resume_text, header_patterns, next_header_patterns)

    if not block_lines:
        lines = [l.strip() for l in resume_text.split("\n") if l.strip()]
        for line in lines:
            if len(line) < 100 and not line.startswith(('•', '*', '-')):
                if re.search(r'developer|engineer|manager|analyst|intern|lead|consultant|associate', line, re.I):
                    block_lines.append(line)

    if not block_lines:
        return []

    exp_entries = []
    current_exp = None

    for line in block_lines:
        clean_l = re.sub(r'^[•\-\*\d\.\s]+', '', line).strip()
        is_role = bool(re.search(r'developer|engineer|manager|analyst|intern|lead|consultant|associate|specialist|architect|administrator', clean_l, re.I))

        if is_role or current_exp is None:
            if current_exp and (current_exp["title"] or current_exp["company"] or current_exp["description"]):
                exp_entries.append(current_exp)
            current_exp = {
                "title": clean_l if clean_l else "Work Experience",
                "company": "",
                "duration": "",
                "description": ""
            }
        else:
            if current_exp:
                date_match = re.search(r'(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec|\d{4})[\s\-–\d\w]*', line, re.I)
                if date_match and not current_exp["duration"]:
                    current_exp["duration"] = line
                elif not current_exp["company"] and len(clean_l) < 60:
                    current_exp["company"] = clean_l
                else:
                    if current_exp["description"]:
                        current_exp["description"] += "\n" + line
                    else:
                        current_exp["description"] = line

    if current_exp and (current_exp["title"] or current_exp["company"] or current_exp["description"]):
        exp_entries.append(current_exp)

    valid_exp = []
    seen_titles = set()
    for e in exp_entries:
        title_clean = e["title"].strip()
        if len(title_clean) > 2 and title_clean.lower() not in seen_titles and title_clean.lower() != "experience":
            seen_titles.add(title_clean.lower())
            valid_exp.append(e)

    return valid_exp


def extract_projects_structured(resume_text, matched_skills=None):
    if matched_skills is None:
        matched_skills = []

    header_patterns = [
        r'^\s*(?:key\s+|academic\s+|personal\s+|technical\s+|major\s+)?projects?\s*(?:work|portfolio)?\s*(?::|-)?\s*$',
        r'^\s*projects?\s*$',
        r'projects?\s+and\s+achievements'
    ]
    next_header_patterns = [
        r'^\s*(?:work\s+)?experience\s*$', r'^\s*employment\s*$', r'^\s*education\s*$',
        r'^\s*skills?\s*$', r'^\s*certificat(?:ions?|es?)\s*$', r'^\s*declarat(?:ion|ions)\s*$',
        r'^\s*achievements?\s*$', r'^\s*languages?\s*$', r'^\s*hobbies\s*$'
    ]

    block_lines = extract_section_block(resume_text, header_patterns, next_header_patterns)

    if not block_lines:
        lines = [l.strip() for l in resume_text.split("\n") if l.strip()]
        for line in lines:
            if re.search(r'\bproject[s]?\b', line, re.I) and len(line) < 120 and not re.search(r'^\s*projects?\s*$', line, re.I):
                block_lines.append(line)

    if not block_lines:
        return []

    projects = []
    current_proj = None

    for line in block_lines:
        is_title = False
        clean_l = re.sub(r'^[•\-\*\d\.\s]+', '', line).strip()

        if line.endswith(":") or re.match(r'^(?:project\s*\d*|title|name)\s*:', line, re.I) or (len(clean_l) < 60 and not clean_l.endswith(".")):
            is_title = True

        techs_found = extract_all_skills_from_text(line)

        if is_title or current_proj is None:
            if current_proj and (current_proj["name"] or current_proj["description"]):
                projects.append(current_proj)
            current_proj = {
                "name": clean_l if clean_l else "Project Entry",
                "description": "",
                "technologies": techs_found if techs_found else list(matched_skills)
            }
        else:
            if current_proj:
                if current_proj["description"]:
                    current_proj["description"] += "\n" + line
                else:
                    current_proj["description"] = line
                for t in techs_found:
                    if t not in current_proj["technologies"]:
                        current_proj["technologies"].append(t)

    if current_proj and (current_proj["name"] or current_proj["description"]):
        projects.append(current_proj)

    valid_projects = []
    seen_names = set()
    for p in projects:
        name_clean = p["name"].strip()
        if len(name_clean) > 2 and name_clean.lower() not in seen_names and name_clean.lower() != "projects":
            seen_names.add(name_clean.lower())
            valid_projects.append(p)

    return valid_projects


# ==========================================
# MAIN ATS MATCHING ENGINE
# ==========================================

def calculate_ats_score(data):
    file_path = data.get("file_path", "")
    job_desc = data.get("job_description", "")
    job_skills = data.get("job_skills", [])
    job_experience = data.get("job_experience", "")

    # 1. Extract Resume Text
    try:
        raw_text = extract_text(file_path)
        if not raw_text or not raw_text.strip():
            return {
                "success": False,
                "error": "Resume parsing failed. The document appears to be empty or unreadable."
            }
    except Exception as e:
        return {
            "success": False,
            "error": f"Resume parsing failed. Reason: {str(e)}"
        }

    # 2. Extract Candidate Info
    candidate_name = extract_name(raw_text)
    email = extract_email(raw_text)
    phone = extract_phone(raw_text)

    # 3. Extract Structured Sections
    structured_education = extract_education_structured(raw_text)
    structured_experience = extract_experience_structured(raw_text)
    structured_projects = extract_projects_structured(raw_text, [])
    structured_certs = extract_certifications_structured(raw_text)
    structured_langs = extract_languages_structured(raw_text)
    structured_achieve = extract_achievements_structured(raw_text)

    # 4. Match Skills (Weight = 40%)
    skill_result = match_skills_strictly(raw_text, job_skills)
    skills_match_score = skill_result["score"]
    matched_skills = skill_result["matched"]
    missing_skills = skill_result["missing"]
    extra_skills = skill_result["extra"]
    all_extracted_skills = matched_skills + extra_skills

    # 5. Match Experience (Weight = 20%)
    experience_match_score = match_experience(raw_text, job_experience, structured_experience)

    # 6. Match Education (Weight = 15%)
    edu_result = match_education(raw_text, "", job_desc, structured_education)
    education_match_score = edu_result["score"]

    # 7. Projects Match (Weight = 15%)
    project_match_score = match_projects(raw_text, structured_projects, job_skills, job_desc)

    # 8. Certifications Match (Weight = 5%)
    certifications_match_score = match_certifications(raw_text, structured_certs, job_desc)

    # 9. Resume Completeness (Weight = 5%)
    completeness_score = calculate_completeness_score(
        candidate_name, email, phone, "Not specified",
        all_extracted_skills, structured_education, structured_experience, structured_projects, structured_certs
    )

    # 10. Weighted Score Calculation
    skills_weight = round(skills_match_score * 0.40)
    exp_weight = round(experience_match_score * 0.20)
    edu_weight = round(education_match_score * 0.15)
    proj_weight = round(project_match_score * 0.15)
    cert_weight = round(certifications_match_score * 0.05)
    comp_weight = round(completeness_score * 0.05)

    ats_score = min(max(skills_weight + exp_weight + edu_weight + proj_weight + cert_weight + comp_weight, 0), 100)

    # 10. Pass / Fail Status
    status_backend = "Accepted" if ats_score >= 60 else "Rejected"
    candidate_status = "Qualified" if ats_score >= 75 else ("Partially Qualified" if ats_score >= 60 else "Not Qualified")

    breakdown = {
        "skills": skills_weight,
        "experience": exp_weight,
        "education": edu_weight,
        "projects": proj_weight,
        "certifications": cert_weight,
        "resumeCompleteness": comp_weight
    }

    return {
        "success": True,
        "atsScore": ats_score,
        "status": candidate_status,
        "breakdown": breakdown,
        "overall_ats_score": ats_score,
        "matchScore": ats_score,
        "matched_skills": matched_skills,
        "missing_skills": missing_skills,
        "extra_skills": extra_skills,
        "extracted_skills": all_extracted_skills,
        "skills": all_extracted_skills,
        "technologies": all_extracted_skills,
        "candidate_name": candidate_name,
        "email": email,
        "phone": phone,
        "education": structured_education,
        "experience": structured_experience,
        "projects": structured_projects,
        "certifications": structured_certs,
        "languages": structured_langs,
        "achievements": structured_achieve,
        "skills_match_score": round(skills_match_score),
        "experience_match_score": round(experience_match_score),
        "education_match_score": round(education_match_score),
        "project_match_score": round(project_match_score),
        "certifications_match_score": round(certifications_match_score),
        "completeness_score": round(completeness_score),
        "status_backend": status_backend,
        "extracted_text": raw_text[:2000],
        "raw_text": raw_text[:2000],
    }


# ==========================================
# CLI ENTRY POINT
# ==========================================

if __name__ == "__main__":
    try:
        if len(sys.argv) > 1 and sys.argv[1] == "--init":
            print("ATS Matcher initialized successfully.", file=sys.stderr)
            sys.exit(0)

        input_data = sys.stdin.read()
        if not input_data.strip():
            print(json.dumps({"success": False, "error": "No input data received"}))
            sys.exit(1)

        data = json.loads(input_data)
        result = calculate_ats_score(data)
        print(json.dumps(result, ensure_ascii=False))

    except Exception as e:
        print(json.dumps({
            "success": False,
            "error": str(e),
            "traceback": traceback.format_exc()
        }))
